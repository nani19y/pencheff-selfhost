"""Generate downloadable reports in DOCX / CSV / JSON / PDF."""
from __future__ import annotations

import asyncio
import csv
import io
import json
import logging
import os
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from ..config import get_settings
from ..db.models import Engagement, Finding, Report, Scan, Target, WorkspaceBranding

log = logging.getLogger(__name__)

SEVERITY_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}
SEVERITY_HEX = {
    "critical": "C00000", "high": "E06666", "medium": "E69138",
    "low": "6FA8DC", "info": "B7B7B7",
}


def _storage_dir() -> Path:
    p = Path(get_settings().report_storage_dir)
    p.mkdir(parents=True, exist_ok=True)
    return p


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def _sort_findings(findings: list[Finding]) -> list[Finding]:
    return sorted(findings, key=lambda f: (SEVERITY_ORDER.get((f.severity or "info").lower(), 99), f.created_at))


def _group_findings(findings: list[Finding]) -> list[Finding]:
    """Collapse findings sharing (title, severity, category) into one card.

    IaC / CI-CD scanners (checkov, zizmor) emit one row per affected file, so a
    single policy — "excessive permissions", "unpinned action" — can appear 20-70
    times. Grouping keeps the report readable without losing the affected-file
    inventory: the representative's ``endpoint`` becomes the merged location list
    and its title is annotated with the occurrence count.

    Safe to mutate in place — ``generate_report`` closes the findings' DB session
    before calling this and only ever commits the Report row afterwards, so these
    edits never persist back to the findings table.
    """
    from collections import OrderedDict

    groups: "OrderedDict[tuple, list[Finding]]" = OrderedDict()
    for f in findings:
        key = ((f.title or "").strip(), (f.severity or "info").lower(), (f.category or "").lower())
        groups.setdefault(key, []).append(f)

    merged: list[Finding] = []
    for members in groups.values():
        rep = members[0]
        if len(members) > 1:
            locs: list[str] = []
            for m in members:
                if m.endpoint and m.endpoint not in locs:
                    locs.append(m.endpoint)
            shown = locs[:12]
            extra = len(locs) - len(shown)
            rep.endpoint = ", ".join(shown) + (f" (+{extra} more)" if extra > 0 else "")
            rep.title = f"{rep.title} ({len(members)}×)"
        merged.append(rep)
    return merged


# ── JSON ──────────────────────────────────────────────────────────────

def write_json(scan: Scan, target: Target, findings: list[Finding], path: Path) -> int:
    data = {
        "report_metadata": {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "target": target.base_url,
            "target_name": target.name,
            "scan_id": scan.id,
            "profile": scan.profile,
            "grade": scan.grade,
            "score": scan.score,
        },
        "summary": scan.summary or {},
        "findings": [
            {
                "id": f.id, "title": f.title, "severity": f.severity,
                "category": f.category, "owasp_category": f.owasp_category,
                "cwe_id": f.cwe_id, "cvss_score": f.cvss_score, "cvss_vector": f.cvss_vector,
                "endpoint": f.endpoint, "parameter": f.parameter,
                "description": f.description, "remediation": f.remediation,
                "evidence": f.evidence, "references": f.references_,
                "verification_status": f.verification_status,
                "suppressed": f.suppressed, "suppress_reason": f.suppress_reason,
                "last_rechecked_at": f.last_rechecked_at.isoformat() if f.last_rechecked_at else None,
            }
            for f in _sort_findings(findings)
        ],
    }
    payload = json.dumps(data, indent=2, default=str)
    path.write_text(payload, encoding="utf-8")
    return len(payload.encode())


# ── CSV ───────────────────────────────────────────────────────────────

def write_csv(scan: Scan, target: Target, findings: list[Finding], path: Path) -> int:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["id", "title", "severity", "category", "owasp", "cwe", "cvss_score",
                "endpoint", "parameter", "verification_status", "suppressed",
                "description", "remediation"])
    for f in _sort_findings(findings):
        w.writerow([
            f.id, f.title, f.severity, f.category, f.owasp_category or "",
            f.cwe_id or "", f.cvss_score or "", f.endpoint or "", f.parameter or "",
            f.verification_status, "yes" if f.suppressed else "no",
            (f.description or "").replace("\n", " ")[:2000],
            (f.remediation or "").replace("\n", " ")[:2000],
        ])
    path.write_text(buf.getvalue(), encoding="utf-8")
    return path.stat().st_size


# ── DOCX ──────────────────────────────────────────────────────────────

def write_docx(scan: Scan, target: Target, findings: list[Finding], path: Path) -> int:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title = doc.add_heading("Pencheff Security Assessment Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run(f"Target: ").bold = True
    meta.add_run(f"{target.name} — {target.base_url}\n")
    meta.add_run(f"Scan ID: ").bold = True
    meta.add_run(f"{scan.id}\n")
    meta.add_run(f"Generated: ").bold = True
    meta.add_run(f"{datetime.now(timezone.utc).isoformat()}\n")
    meta.add_run(f"Grade: ").bold = True
    grade_run = meta.add_run(f"{scan.grade or '-'}  (score: {scan.score or 0}/100)\n")
    grade_run.bold = True
    grade_run.font.size = Pt(16)

    doc.add_heading("Executive Summary", level=1)
    summary = scan.summary or {}
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Shading Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Critical"
    hdr[1].text = "High"
    hdr[2].text = "Medium"
    hdr[3].text = "Low"
    hdr[4].text = "Info"
    row = table.add_row().cells
    row[0].text = str(summary.get("critical", 0))
    row[1].text = str(summary.get("high", 0))
    row[2].text = str(summary.get("medium", 0))
    row[3].text = str(summary.get("low", 0))
    row[4].text = str(summary.get("info", 0))

    doc.add_heading("Findings", level=1)
    for f in _sort_findings(findings):
        h = doc.add_heading(f"[{f.severity.upper()}] {f.title}", level=2)
        try:
            color = SEVERITY_HEX.get((f.severity or "info").lower(), "000000")
            for run in h.runs:
                run.font.color.rgb = RGBColor.from_string(color)
        except Exception:
            pass
        p = doc.add_paragraph()
        p.add_run("Category: ").bold = True
        p.add_run(f"{f.category}   ")
        if f.owasp_category:
            p.add_run("OWASP: ").bold = True
            p.add_run(f"{f.owasp_category}   ")
        if f.cvss_score:
            p.add_run("CVSS: ").bold = True
            p.add_run(f"{f.cvss_score} ({f.cvss_vector or ''})   ")
        if f.endpoint:
            p.add_run("Endpoint: ").bold = True
            p.add_run(f"{f.endpoint}   ")
        if f.parameter:
            p.add_run("Parameter: ").bold = True
            p.add_run(f"{f.parameter}")
        if f.description:
            doc.add_paragraph().add_run("Description: ").bold = True
            doc.add_paragraph(f.description)
        if f.remediation:
            doc.add_paragraph().add_run("Remediation: ").bold = True
            doc.add_paragraph(f.remediation)
        doc.add_paragraph("")

    doc.add_heading("Compliance Mapping", level=1)
    doc.add_paragraph(
        "Findings in this report map to controls from OWASP Top 10 2021, SOC 2 (CC6/CC7), "
        "PCI-DSS 4.0, NIST 800-53, ISO 27001:2022, and HIPAA Security Rule. "
        "This report is suitable for attachment to SOC 2 Type II and similar audit evidence packages."
    )
    notices = doc.add_paragraph()
    notices.add_run(
        "OWASP and OWASP Top 10 are trademarks of the OWASP Foundation. "
        "Pencheff is not affiliated with or endorsed by OWASP. Other framework "
        "names referenced above are the property of their respective owners and "
        "are used for identification only."
    ).italic = True

    doc.save(str(path))
    return path.stat().st_size


# ── PDF ───────────────────────────────────────────────────────────────

def write_pdf(scan: Scan, target: Target, findings: list[Finding], path: Path) -> int:
    from weasyprint import HTML

    rows = []
    for f in _sort_findings(findings):
        sev = (f.severity or "info").lower()
        color = "#" + SEVERITY_HEX.get(sev, "888888")
        rows.append(f"""
        <div class="finding sev-{sev}">
          <h3 style="color:{color}">[{sev.upper()}] {(f.title or '').replace('<', '&lt;')}</h3>
          <p><b>Category:</b> {f.category}
             {('<b>OWASP:</b> ' + f.owasp_category) if f.owasp_category else ''}
             {('<b>CVSS:</b> ' + str(f.cvss_score)) if f.cvss_score else ''}</p>
          <p><b>Endpoint:</b> {f.endpoint or '-'} &nbsp; <b>Parameter:</b> {f.parameter or '-'}</p>
          <p><b>Description:</b> {(f.description or '').replace('<', '&lt;')}</p>
          <p><b>Remediation:</b> {(f.remediation or '').replace('<', '&lt;')}</p>
          <p><b>Status:</b> {f.verification_status}{' (suppressed)' if f.suppressed else ''}</p>
        </div>
        """)
    summary = scan.summary or {}
    html = f"""
    <!doctype html>
    <html><head><style>
      body {{ font-family: Helvetica, Arial, sans-serif; color:#111; }}
      h1 {{ border-bottom: 4px solid #000; padding-bottom: 8px; }}
      .grade {{ display:inline-block; padding: 10px 30px; border:3px solid #000;
               font-size:48px; font-weight:900; background:#FFD23F; box-shadow:6px 6px 0 #000; }}
      table {{ border-collapse:collapse; margin: 16px 0; }}
      th, td {{ border:2px solid #000; padding:6px 12px; }}
      .finding {{ margin:16px 0; padding:12px; border:3px solid #000;
                 box-shadow:4px 4px 0 #000; background:#FDFBF5; }}
    </style></head><body>
      <h1>Pencheff Security Assessment</h1>
      <p><b>Target:</b> {target.name} — {target.base_url}<br>
         <b>Scan ID:</b> {scan.id}<br>
         <b>Generated:</b> {datetime.now(timezone.utc).isoformat()}</p>
      <p class="grade">{scan.grade or '-'}</p>
      <p><b>Score:</b> {scan.score or 0}/100</p>
      <h2>Summary</h2>
      <table>
        <tr><th>Critical</th><th>High</th><th>Medium</th><th>Low</th><th>Info</th></tr>
        <tr>
          <td>{summary.get('critical', 0)}</td>
          <td>{summary.get('high', 0)}</td>
          <td>{summary.get('medium', 0)}</td>
          <td>{summary.get('low', 0)}</td>
          <td>{summary.get('info', 0)}</td>
        </tr>
      </table>
      <h2>Findings</h2>
      {''.join(rows)}
      <hr style="margin-top:32px;border:none;border-top:1px solid #888;">
      <p style="font-size:10px;color:#555;font-style:italic;">
        OWASP and OWASP Top 10 are trademarks of the OWASP Foundation.
        Pencheff is not affiliated with or endorsed by OWASP. Other framework
        names referenced in this report are the property of their respective
        owners and are used for identification only.
      </p>
    </body></html>
    """
    HTML(string=html).write_pdf(str(path))
    return path.stat().st_size


# ── Markdown export (consultancy-friendly) ────────────────────────────

def write_markdown(scan: Scan, target: Target, findings: list[Finding], path: Path,
                   branding: WorkspaceBranding | None = None,
                   threat_model: dict | None = None) -> int:
    lines: list[str] = []
    if branding and branding.logo_url:
        lines.append(f"![logo]({branding.logo_url})\n")
    lines.append(f"# Pencheff Security Assessment — {target.name}\n")
    if branding and branding.opening_letter_md:
        lines.append(branding.opening_letter_md.strip() + "\n")
    lines.append(f"- **Target**: {target.base_url}")
    lines.append(f"- **Scan ID**: `{scan.id}`")
    lines.append(f"- **Profile**: `{scan.profile}`")
    if scan.grade:
        lines.append(f"- **Grade**: **{scan.grade}** ({scan.score or 0}/100)")
    lines.append(f"- **Generated**: {datetime.now(timezone.utc).isoformat()}\n")
    if branding and branding.methodology_md:
        lines.append("## Methodology\n")
        lines.append(branding.methodology_md.strip() + "\n")
    summary = scan.summary or {}
    if threat_model:
        from .threat_model import render_markdown as _render_threat_model
        # Render the engagement-scoped STRIDE/DREAD model inline. The
        # service emits its own ``# STRIDE/DREAD Threat Model`` heading,
        # so we wrap it in a section anchor for the renderer's TOC.
        lines.append("## Threat model")
        lines.append("")
        lines.append(_render_threat_model(threat_model).strip())
        lines.append("")
    lines.append("## Summary\n")
    lines.append("| Critical | High | Medium | Low | Info |")
    lines.append("|---|---|---|---|---|")
    lines.append(
        f"| {summary.get('critical', 0)} | {summary.get('high', 0)} | "
        f"{summary.get('medium', 0)} | {summary.get('low', 0)} | "
        f"{summary.get('info', 0)} |\n"
    )
    lines.append("## Findings\n")
    for f in _sort_findings(findings):
        lines.append(f"### {f.severity.upper()} — {f.title}")
        lines.append(f"- **OWASP**: {f.owasp_category or '-'}  **CWE**: {f.cwe_id or '-'}")
        if f.cvss_score:
            lines.append(f"- **CVSS**: {f.cvss_score}")
        lines.append(f"- **Endpoint**: `{f.endpoint or '-'}`  **Parameter**: `{f.parameter or '-'}`")
        lines.append(f"- **Status**: {f.verification_status}")
        lines.append("")
        if f.description:
            lines.append(f.description.strip() + "\n")
        if f.remediation:
            lines.append(f"**Remediation**: {f.remediation.strip()}\n")
    if branding and branding.footer_text:
        lines.append("---\n")
        lines.append(branding.footer_text)
    payload = "\n".join(lines)
    path.write_text(payload, encoding="utf-8")
    return len(payload.encode())


# ── Delta report ──────────────────────────────────────────────────────

def write_delta_markdown(
    scan_a: Scan, scan_b: Scan, target: Target,
    findings_a: list[Finding], findings_b: list[Finding], path: Path,
    branding: WorkspaceBranding | None = None,
) -> int:
    """Emit a Markdown delta report comparing scan_a (older) to scan_b (newer).

    A finding is keyed by (title, endpoint, parameter, owasp_category) for
    the diff. New = present in B not A. Fixed = present in A not B.
    Regressed = present in both but severity worsened.
    """
    def _key(f: Finding) -> tuple:
        return (f.title, f.endpoint or "", f.parameter or "", f.owasp_category or "")

    a_index = {_key(f): f for f in findings_a}
    b_index = {_key(f): f for f in findings_b}
    new = [f for k, f in b_index.items() if k not in a_index]
    fixed = [f for k, f in a_index.items() if k not in b_index]
    regressed = [
        f for k, f in b_index.items()
        if k in a_index and SEVERITY_ORDER.get(f.severity, 99) < SEVERITY_ORDER.get(a_index[k].severity, 99)
    ]

    lines: list[str] = []
    if branding and branding.logo_url:
        lines.append(f"![logo]({branding.logo_url})\n")
    lines.append(f"# Re-test report — {target.name}\n")
    lines.append(f"Comparing scan `{scan_a.id[:8]}` ({scan_a.created_at.isoformat()}) → "
                 f"`{scan_b.id[:8]}` ({scan_b.created_at.isoformat()}).\n")
    lines.append(f"- **New findings**: {len(new)}")
    lines.append(f"- **Fixed findings**: {len(fixed)}")
    lines.append(f"- **Regressed (severity ↑)**: {len(regressed)}\n")

    def _section(title: str, rows: list[Finding]):
        lines.append(f"## {title}\n")
        if not rows:
            lines.append("_none_\n")
            return
        for f in _sort_findings(rows):
            lines.append(f"- **{f.severity.upper()}** — {f.title} (`{f.endpoint or '-'}`)")
        lines.append("")

    _section("New", new)
    _section("Fixed", fixed)
    _section("Regressed", regressed)
    payload = "\n".join(lines)
    path.write_text(payload, encoding="utf-8")
    return len(payload.encode())


# ── Load-test reports (profile="load") ────────────────────────────────
# Load scans have no findings — they render the load_report on scan.summary
# (stages, breaking point, SLA thresholds, status buckets, agent verdict).


def _load_ctx(scan: Scan, target: Target) -> dict:
    r = (scan.summary or {}).get("load_report") or {}
    extra = r.get("extra") or {}
    return {
        "kind": r.get("kind", target.kind),
        "requests": r.get("total_requests", 0),
        "rps": r.get("rps", 0),
        "duration": r.get("duration_s", 0),
        "error_rate": r.get("error_rate", 0),            # raw non-2xx
        "load_error_rate": extra.get("load_error_rate", 0),  # 5xx+429+timeouts
        "latency": r.get("latency_ms") or {},
        "stages": r.get("stages") or [],
        "thresholds": r.get("threshold_results") or [],
        "breaking_point": r.get("breaking_point"),
        "status_buckets": extra.get("status_buckets") or {},
        "narrative": r.get("agent_narrative") or "",
        "findings": extra.get("findings") or [],
        "personas": extra.get("personas") or [],
        "profile_shape": extra.get("profile_shape") or "",
        "cost_cents": r.get("cost_cents", 0),
    }


def _pct(v) -> str:
    try:
        return f"{float(v) * 100:.1f}%"
    except (TypeError, ValueError):
        return "-"


def write_load_json(scan: Scan, target: Target, path: Path) -> int:
    data = {
        "report_type": "load_test",
        "target": {"name": target.name, "base_url": str(target.base_url),
                   "kind": target.kind},
        "scan": {"id": scan.id, "profile": scan.profile, "grade": scan.grade,
                 "generated_at": datetime.now(timezone.utc).isoformat()},
        "load_report": (scan.summary or {}).get("load_report") or {},
    }
    payload = json.dumps(data, indent=2, default=str)
    path.write_text(payload, encoding="utf-8")
    return len(payload.encode())


def write_load_csv(scan: Scan, target: Target, path: Path) -> int:
    ctx = _load_ctx(scan, target)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["stage", "vus", "target_rps", "requests", "rps", "p95_ms",
                    "load_error_rate", "non_2xx_rate", "sla_hold"])
        for i, st in enumerate(ctx["stages"], 1):
            w.writerow([i, st.get("vus"), st.get("target_rps"),
                        st.get("total_requests"), st.get("rps"),
                        st.get("p95_ms"), st.get("load_error_rate"),
                        st.get("error_rate"), st.get("sla_hold")])
    return path.stat().st_size


def write_load_docx(scan: Scan, target: Target, path: Path) -> int:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    ctx = _load_ctx(scan, target)
    doc = Document()
    title = doc.add_heading("Pencheff Load-Test Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run("Target: ").bold = True
    meta.add_run(f"{target.name} — {target.base_url}\n")
    meta.add_run("Scan ID: ").bold = True
    meta.add_run(f"{scan.id}\n")
    meta.add_run("Generated: ").bold = True
    meta.add_run(f"{datetime.now(timezone.utc).isoformat()}\n")
    meta.add_run("Grade: ").bold = True
    g = meta.add_run(f"{scan.grade or '-'}\n")
    g.bold = True
    g.font.size = Pt(16)

    doc.add_heading("Summary", level=1)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Light Shading Accent 1"
    hdr = t.rows[0].cells
    for i, h in enumerate(
        ["Requests", "Throughput", "Load errors", "Non-2xx", "p95"]
    ):
        hdr[i].text = h
    row = t.add_row().cells
    row[0].text = str(ctx["requests"])
    row[1].text = f"{ctx['rps']} rps"
    row[2].text = _pct(ctx["load_error_rate"])
    row[3].text = _pct(ctx["error_rate"])
    row[4].text = f"{round(ctx['latency'].get('p95', 0))}ms"

    bp = ctx["breaking_point"]
    doc.add_heading("Breaking point", level=1)
    if bp:
        doc.add_paragraph(
            f"~{bp.get('vus')} VUs — load errors "
            f"{_pct(bp.get('load_error_rate', bp.get('error_rate')))}, "
            f"p95 {round(bp.get('p95_ms') or 0)}ms."
        )
    else:
        doc.add_paragraph("Held SLA through the full ladder — no breaking point.")

    if ctx["thresholds"]:
        doc.add_heading("SLA thresholds", level=1)
        for th in ctx["thresholds"]:
            doc.add_paragraph(
                f"{th.get('metric')}: actual {th.get('actual')} / limit "
                f"{th.get('limit')} — {'PASS' if th.get('pass') else 'FAIL'}",
                style="List Bullet",
            )

    if ctx["stages"]:
        doc.add_heading("Stages", level=1)
        st = doc.add_table(rows=1, cols=5)
        st.style = "Light Shading Accent 1"
        h = st.rows[0].cells
        for i, lbl in enumerate(["VUs", "RPS", "p95", "Load err", "SLA"]):
            h[i].text = lbl
        for stg in ctx["stages"]:
            r = st.add_row().cells
            r[0].text = str(stg.get("vus"))
            r[1].text = str(round(stg.get("rps", 0), 1))
            r[2].text = f"{round(stg.get('p95_ms', 0))}ms"
            r[3].text = _pct(stg.get("load_error_rate"))
            r[4].text = "PASS" if stg.get("sla_hold") else "BREACH"

    sb = ctx["status_buckets"]
    if sb:
        doc.add_heading("Status codes", level=1)
        doc.add_paragraph(
            ", ".join(f"{k}: {v}" for k, v in sb.items() if v)
        )

    if ctx["narrative"]:
        doc.add_heading("Performance-engineer verdict", level=1)
        doc.add_paragraph(ctx["narrative"])

    if ctx["findings"]:
        doc.add_heading("Findings", level=1)
        for f in ctx["findings"]:
            doc.add_paragraph(
                f"[{str(f.get('severity', 'note')).upper()}] {f.get('title', '')}"
                f" — {f.get('detail', '')}",
                style="List Bullet",
            )

    doc.save(str(path))
    return path.stat().st_size


def write_load_pdf(scan: Scan, target: Target, path: Path) -> int:
    from weasyprint import HTML

    ctx = _load_ctx(scan, target)

    def esc(x) -> str:
        return str(x).replace("<", "&lt;")

    bp = ctx["breaking_point"]
    bp_html = (
        f"~{bp.get('vus')} VUs — load errors "
        f"{_pct(bp.get('load_error_rate', bp.get('error_rate')))}, "
        f"p95 {round(bp.get('p95_ms') or 0)}ms"
        if bp else "Held SLA through the full ladder — no breaking point."
    )
    stage_rows = "".join(
        f"<tr><td>{s.get('vus')}</td><td>{round(s.get('rps', 0), 1)}</td>"
        f"<td>{round(s.get('p95_ms', 0))}ms</td>"
        f"<td>{_pct(s.get('load_error_rate'))}</td>"
        f"<td>{'PASS' if s.get('sla_hold') else 'BREACH'}</td></tr>"
        for s in ctx["stages"]
    )
    thr_rows = "".join(
        f"<li>{esc(t.get('metric'))}: actual {esc(t.get('actual'))} / limit "
        f"{esc(t.get('limit'))} — <b>{'PASS' if t.get('pass') else 'FAIL'}</b></li>"
        for t in ctx["thresholds"]
    )
    sb = ", ".join(f"{k}: {v}" for k, v in ctx["status_buckets"].items() if v)
    findings_html = "".join(
        f"<li><b>[{esc(str(f.get('severity', 'note')).upper())}]</b> "
        f"{esc(f.get('title', ''))} — {esc(f.get('detail', ''))}</li>"
        for f in ctx["findings"]
    )
    html = f"""
    <!doctype html><html><head><style>
      body {{ font-family: Helvetica, Arial, sans-serif; color:#111; }}
      h1 {{ border-bottom: 4px solid #000; padding-bottom: 8px; }}
      .grade {{ display:inline-block; padding: 10px 30px; border:3px solid #000;
               font-size:48px; font-weight:900; background:#FFD23F; box-shadow:6px 6px 0 #000; }}
      table {{ border-collapse:collapse; margin: 16px 0; }}
      th, td {{ border:2px solid #000; padding:6px 12px; }}
      .box {{ margin:16px 0; padding:12px; border:3px solid #000;
             box-shadow:4px 4px 0 #000; background:#FDFBF5; }}
    </style></head><body>
      <h1>Pencheff Load-Test Report</h1>
      <p><b>Target:</b> {esc(target.name)} — {esc(target.base_url)}<br>
         <b>Scan ID:</b> {esc(scan.id)}<br>
         <b>Generated:</b> {datetime.now(timezone.utc).isoformat()}</p>
      <p class="grade">{scan.grade or '-'}</p>
      <h2>Summary</h2>
      <table>
        <tr><th>Requests</th><th>Throughput</th><th>Load errors</th><th>Non-2xx</th><th>p95</th></tr>
        <tr><td>{ctx['requests']}</td><td>{ctx['rps']} rps</td>
            <td>{_pct(ctx['load_error_rate'])}</td><td>{_pct(ctx['error_rate'])}</td>
            <td>{round(ctx['latency'].get('p95', 0))}ms</td></tr>
      </table>
      <div class="box"><b>Breaking point:</b> {bp_html}</div>
      {f'<h2>SLA thresholds</h2><ul>{thr_rows}</ul>' if thr_rows else ''}
      {f'<h2>Stages</h2><table><tr><th>VUs</th><th>RPS</th><th>p95</th><th>Load err</th><th>SLA</th></tr>{stage_rows}</table>' if stage_rows else ''}
      {f'<p><b>Status codes:</b> {esc(sb)}</p>' if sb else ''}
      {f'<h2>Performance-engineer verdict</h2><p>{esc(ctx["narrative"])}</p>' if ctx['narrative'] else ''}
      {f'<h2>Findings</h2><ul>{findings_html}</ul>' if findings_html else ''}
    </body></html>
    """
    HTML(string=html).write_pdf(str(path))
    return path.stat().st_size


def write_load_markdown(scan: Scan, target: Target, path: Path) -> int:
    ctx = _load_ctx(scan, target)
    L = [f"# Pencheff Load-Test Report — {target.name}", ""]
    L.append(f"- **Target**: {target.base_url}")
    L.append(f"- **Scan ID**: `{scan.id}`")
    L.append(f"- **Grade**: {scan.grade or '-'}")
    L.append(f"- **Generated**: {datetime.now(timezone.utc).isoformat()}\n")
    L.append("## Summary")
    L.append(f"- Requests: {ctx['requests']} · Throughput: {ctx['rps']} rps")
    L.append(f"- Load errors (5xx/429/timeout): {_pct(ctx['load_error_rate'])}")
    L.append(f"- Non-2xx (raw): {_pct(ctx['error_rate'])}")
    L.append(f"- p95 latency: {round(ctx['latency'].get('p95', 0))}ms\n")
    bp = ctx["breaking_point"]
    L.append("## Breaking point")
    L.append(
        f"~{bp.get('vus')} VUs — load errors "
        f"{_pct(bp.get('load_error_rate', bp.get('error_rate')))}, "
        f"p95 {round(bp.get('p95_ms') or 0)}ms\n"
        if bp else "Held SLA through the full ladder — no breaking point.\n"
    )
    if ctx["stages"]:
        L.append("## Stages")
        L.append("| VUs | RPS | p95 | Load err | SLA |")
        L.append("|----|----|----|----|----|")
        for s in ctx["stages"]:
            L.append(
                f"| {s.get('vus')} | {round(s.get('rps', 0), 1)} | "
                f"{round(s.get('p95_ms', 0))}ms | {_pct(s.get('load_error_rate'))} | "
                f"{'PASS' if s.get('sla_hold') else 'BREACH'} |"
            )
        L.append("")
    if ctx["status_buckets"]:
        L.append("## Status codes")
        L.append(", ".join(f"{k}: {v}" for k, v in ctx["status_buckets"].items() if v) + "\n")
    if ctx["narrative"]:
        L.append("## Performance-engineer verdict")
        L.append(ctx["narrative"] + "\n")
    if ctx["findings"]:
        L.append("## Findings")
        for f in ctx["findings"]:
            L.append(
                f"- **[{str(f.get('severity', 'note')).upper()}]** "
                f"{f.get('title', '')} — {f.get('detail', '')}"
            )
    payload = "\n".join(L)
    path.write_text(payload, encoding="utf-8")
    return len(payload.encode())


LOAD_WRITERS = {
    "json": write_load_json, "csv": write_load_csv, "docx": write_load_docx,
    "pdf": write_load_pdf, "markdown": write_load_markdown,
}


# ── Orchestrator ──────────────────────────────────────────────────────

WRITERS = {"json": write_json, "csv": write_csv, "docx": write_docx, "pdf": write_pdf, "markdown": write_markdown}
EXT = {"json": "json", "csv": "csv", "docx": "docx", "pdf": "pdf", "markdown": "md"}


async def generate_report(report_id: str) -> None:
    settings = get_settings()
    engine = create_async_engine(settings.database_url, future=True)
    Session = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)

    async with Session() as db:
        report = (await db.execute(select(Report).where(Report.id == report_id))).scalar_one_or_none()
        if not report:
            return
        scan = (await db.execute(select(Scan).where(Scan.id == report.scan_id))).scalar_one()
        target = (await db.execute(select(Target).where(Target.id == scan.target_id))).scalar_one()
        findings = (await db.execute(
            select(Finding).where(Finding.scan_id == scan.id).where(Finding.suppressed.is_(False))
        )).scalars().all()
        branding = (await db.execute(
            select(WorkspaceBranding).where(WorkspaceBranding.workspace_id == scan.workspace_id)
        )).scalar_one_or_none()
        # If the scan was scoped to an engagement that has a threat model
        # attached, pull it so write_markdown can include a "Threat model"
        # section. Reports for ad-hoc scans (no engagement) skip this.
        threat_model: dict | None = None
        if scan.engagement_id:
            eng = (await db.execute(
                select(Engagement).where(Engagement.id == scan.engagement_id)
            )).scalar_one_or_none()
            if eng is not None and eng.threat_model:
                threat_model = eng.threat_model
        compared = None
        compared_findings: list[Finding] = []
        if report.kind == "delta" and report.compared_scan_id:
            compared = (await db.execute(select(Scan).where(Scan.id == report.compared_scan_id))).scalar_one_or_none()
            if compared is not None:
                compared_findings = (await db.execute(
                    select(Finding).where(Finding.scan_id == compared.id).where(Finding.suppressed.is_(False))
                )).scalars().all()

    # Collapse duplicate IaC / CI-CD findings (one row per file) into a single
    # card each. Skipped for delta reports, which diff findings by identity.
    if report.kind != "delta":
        findings = _group_findings(findings)

    fmt = report.format
    # Load scans have no findings — render the load_report instead.
    is_load = scan.profile == "load" or bool((scan.summary or {}).get("load_report"))
    writer = (LOAD_WRITERS if is_load else WRITERS).get(fmt)
    if not writer and report.kind != "delta":
        async with Session() as db:
            r = (await db.execute(select(Report).where(Report.id == report_id))).scalar_one()
            r.status = "failed"
            await db.commit()
        return

    out_dir = _storage_dir() / scan.id
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = f"pencheff_{scan.id[:8]}_{_timestamp()}.{EXT.get(fmt, 'md')}"
    path = out_dir / filename

    try:
        if is_load:
            size = writer(scan, target, path)
        elif report.kind == "delta" and compared is not None:
            size = write_delta_markdown(
                compared, scan, target, compared_findings, findings, path, branding,
            )
        elif fmt == "markdown":
            size = write_markdown(
                scan, target, findings, path, branding, threat_model=threat_model
            )
        else:
            size = writer(scan, target, findings, path)
    except Exception as e:
        log.exception("report generation failed")
        async with Session() as db:
            r = (await db.execute(select(Report).where(Report.id == report_id))).scalar_one()
            r.status = "failed"
            await db.commit()
        return

    async with Session() as db:
        r = (await db.execute(select(Report).where(Report.id == report_id))).scalar_one()
        r.status = "ready"
        r.storage_path = str(path)
        r.bytes = size
        r.generated_at = datetime.now(timezone.utc)
        await db.commit()


def generate_report_sync(report_id: str) -> None:
    asyncio.run(generate_report(report_id))
